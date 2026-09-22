"""内容页 —— 当前 deck 的正文。

本文件是一份**功能完备的最小示例**：管线支持的每一种版式/能力都在这里出现至少
一次，所以它同时是 cookbook —— 想加一页就找个同类函数照抄。

放什么、放哪儿：
  * 版式（进度条、校徽、页脚、标题占位符）在 build_pptx.py，不在这里。
  * 本文件只放**页上**的内容。版式上的文字在普通视图里点不中、改不了，
    所以凡是作者可能要改的字符串（封面标题、章节眉标、目录条目）都写在这里。

生成式输入（由 build/ 下的 prep 脚本产出，本模块只读）：
  assets/figures/manifest.json    prep_figures.py    插图（含矢量 PDF 光栅化）
  assets/formulas/manifest.json   make_formulas.py   公式 / 伪代码
  _refs.json                      prep_refs.py       参考文献（可选）
"""
import json
import os

from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from pptx.util import Emu, Pt

import _pptx_kit as K
from build_pptx import DECK, BAR_TOP, BAR_W, BAR_H, BAR_PITCH, LATIN, EA

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))

M = 53
BODY_W = 1174
CONTENT_Y = 254
COL_W = (BODY_W - 29) / 2          # two-column body, 29px gutter
COL2_X = M + COL_W + 29
CARD_W = (BODY_W - 40) / 3         # three cards, 20px gutters
FOOT_RULE_Y = 660                  # body must stop above this

GOLD_SUB_H = 3
GOLD_GAP = 3                       # gap between the segment and its gold bar

# --- formula scale -----------------------------------------------------------
# Every formula PNG comes out of make_formulas.py at the same DPI from the same
# 10pt base, so ONE PNG PIXEL IS A FIXED PHYSICAL SIZE. Placing every formula at
# the same scale factor is therefore what makes the maths read as one document.
# Scaling each image to fill its own box instead -- the obvious thing to do --
# makes a short formula huge and a long one tiny on the SAME slide. EQ_K is the
# deck standard; a formula is only ever scaled below it when its box is too small.
EQ_K = 0.40                        # ~16pt effective maths size
WALL_H = 386                       # body height a full-page equation wall gets

# DrawingML baseline shifts, in 1/1000 percent.
SUP = 30000
SUB = -25000

# "No Style, No Grid" — lets us draw the booktabs-style rules ourselves.
NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def _load(sub, name="manifest.json"):
    with open(os.path.join(ROOT, "assets", sub, name), encoding="utf-8") as fh:
        return json.load(fh)


FIG = _load("figures")
EQ = _load("formulas")
# Optional: prep_refs.py writes this. The demo inlines its own short list, so a
# missing file must not break the build.
try:
    with open(os.path.join(HERE, "_refs.json"), encoding="utf-8") as _fh:
        REFS = json.load(_fh)
except OSError:
    REFS = []


# ---------------------------------------------------------------------------
# primitives
# ---------------------------------------------------------------------------

def add(slide, el):
    slide.shapes._spTree.append(el)


def tb(slide, tok, x, y, w, h, paragraphs, name="tb", anchor="t"):
    add(slide, K.textbox(K.px(x), K.px(y), K.px(w), K.px(h), paragraphs,
                         name=name, anchor=anchor))


def rule(slide, tok, x, y, w, h=1, key="rule"):
    add(slide, K.rect(K.c(tok, key), K.px(x), K.px(y), K.px(w), K.px(h),
                      name="rule"))


def _c(tok, key):
    """Shorthand colour keys used throughout the deck's body text."""
    return {"k": K.c(tok, "primary-dark"), "n": K.c(tok, "secondary-dark"),
            "m": K.c(tok, "ink"), "g": K.c(tok, "primary"),
            "a": K.c(tok, "alert"), "u": K.c(tok, "ink-muted")}.get(key, key)


def R(tok, items):
    """[(text, colour-key-or-hex, bold)] -> run tuples."""
    return [(t, _c(tok, c), b) for t, c, b in items]


def MT(tok, items):
    """Like R(), but each string is passed through mr() so that U_a(k) and
    UE_a^{s,e}(t) render with real sub/superscripts instead of literal
    underscores. Plain prose is unaffected (mr returns a single run)."""
    out = []
    for t, c, b in items:
        out += mr(t, _c(tok, c), b)
    return out


def P(runs, size=14, ls=1.45, align="l", **kw):
    """A paragraph dict with the boilerplate filled in."""
    d = {"runs": runs, "size": size, "latin": LATIN, "ea": EA,
         "align": align, "line_spacing": ls}
    d.update(kw)
    return d


def mr(expr, color, bold=False):
    """Notation syntax -> runs with real super/subscripts.

    ``UE_a^{s,e}(t)`` becomes UE, sub a, sup s,e, then (t). After ``_`` or ``^``
    either a braced group or a single character is consumed.

    **This consumes every bare ``_`` / ``^``.** That is fine for notation but
    wrong for prose and identifiers: ``my_file`` would render as "my" + subscript
    "f" + "ile", and a URL or a filename would be mangled the same way. So:
      * plain text -> use ``R()``, not this;
      * a literal underscore inside notation -> escape it as ``\\_``.

    Keeping notation as text (rather than an image) means it stays editable and
    searchable.
    """
    out, buf, i = [], "", 0
    while i < len(expr):
        ch = expr[i]
        if ch == "\\" and i + 1 < len(expr) and expr[i + 1] in "_^\\":
            buf += expr[i + 1]           # \_ \^ \\ -> the literal character
            i += 2
            continue
        if ch in "_^":
            if buf:
                out.append((buf, color, bold))
                buf = ""
            j = i + 1
            if j < len(expr) and expr[j] == "{":
                k = expr.find("}", j)
                if k < 0:
                    buf += ch
                    i += 1
                    continue
                tok_, i = expr[j + 1:k], k + 1
            else:
                tok_, i = (expr[j] if j < len(expr) else ""), j + 1
            if tok_:
                out.append((tok_, color, bold, SUB if ch == "_" else SUP))
        else:
            buf += ch
            i += 1
    if buf:
        out.append((buf, color, bold))
    return out


def bullets(slide, tok, x, y, w, h, items, name="bullets", size=16, ls=1.45,
            spc_after=8, indent=171450):
    """items: list of run-lists (see R) — one bulleted paragraph each."""
    ps = []
    for runs in items:
        p = P(runs, size=size, ls=ls, spc_after=spc_after)
        if indent:
            p["marL"] = p["indent"] = 0
            p["marL"], p["indent"] = indent, -indent
            p["bullet"] = K.c(tok, "primary")
        ps.append(p)
    tb(slide, tok, x, y, w, h, ps, name=name)


def label(slide, tok, x, y, w, text, size=14, color="k", h=22):
    tb(slide, tok, x, y, w, h, [P(R(tok, [(text, color, True)]), size=size)],
       name="小标题")


def note(slide, tok, x, y, w, text, size=11, h=20, align="l"):
    tb(slide, tok, x, y, w, h,
       [P(R(tok, [(text, "u", False)]), size=size, align=align, ls=1.3)],
       name="注释")


# tbl() takes a `note=` parameter, which would shadow the function above inside
# its body. Keep a private alias for that call site.
_note_fn = note


# ---------------------------------------------------------------------------
# images and formulas
# ---------------------------------------------------------------------------

def _fit(info, mw, mh, cap=None):
    """Largest scale that fits mw x mh, never above `cap`."""
    s = min(mw / info["w"], mh / info["h"])
    if cap is not None:
        s = min(s, cap)
    return info["w"] * s, info["h"] * s


def _put(slide, info, x, y, w, h, name):
    _, rId = slide.part.get_or_add_image_part(os.path.join(ROOT, info["file"]))
    add(slide, K.picture(rId, K.px(x), K.px(y), K.px(w), K.px(h), name))


def place(slide, info, x, y, mw, mh, align="ctr", valign="t", name="pic", cap=None):
    w, h = _fit(info, mw, mh, cap)
    px = x + (mw - w) / 2 if align == "ctr" else (x + mw - w if align == "r" else x)
    py = y + (mh - h) / 2 if valign == "ctr" else (y + mh - h if valign == "b" else y)
    _put(slide, info, px, py, w, h, name)
    return py + h


def img(slide, tok, key, x, y, mw, mh, align="ctr", valign="t", name=None):
    """A figure from assets/figures, scaled to fit mw x mh without distortion.

    Deliberately NOT capped: a photo or a chart is meant to fill its frame, and
    unlike maths there is no "correct" size to stay consistent with.
    """
    return place(slide, FIG[key], x, y, mw, mh, align, valign, name or key)


def eqk(keys, w, cols=1, gap_x=24, gap_y=10, mh=None):
    """The uniform scale eqgrid() would pick.

    Lets a companion formula placed with ``eqimg(..., k=eqk(...))`` come out at
    exactly the same size as the grid beside it.
    """
    per = (len(keys) + cols - 1) // cols
    col_w = (w - gap_x * (cols - 1)) / cols
    k = min(EQ_K, min(col_w / EQ[x]["w"] for x in keys))
    if mh is not None:
        # Solve for k per column rather than scaling a first estimate: the
        # inter-equation gaps are in canvas px and do NOT shrink with k, so
        # multiplying by mh/tallest would leave the column taller than mh.
        for c in range(cols):
            chunk = keys[c * per:(c + 1) * per]
            room = mh - gap_y * max(0, len(chunk) - 1)
            if room <= 0:
                return 0.01
            k = min(k, room / sum(EQ[x]["h"] for x in chunk))
    return k


def eqplace(slide, tok, slots, k=None):
    """Place formulas into explicitly-sized boxes at ONE shared scale.

    ``slots`` is ``[(key, x, y, w, h)]``. Use this when a slide's formulas sit in
    differently-shaped boxes — a wide objective above a narrow constraint block —
    because fitting each box separately would give them different maths sizes.
    """
    k = EQ_K if k is None else k
    k = min([k] + [w / EQ[key]["w"] for key, _, _, w, _ in slots])
    k = min([k] + [h / EQ[key]["h"] for key, _, _, _, h in slots])
    for key, x, y, w, h in slots:
        info = EQ[key]
        iw, ih = info["w"] * k, info["h"] * k
        _put(slide, info, x + (w - iw) / 2, y + (h - ih) / 2, iw, ih, key)
    return k


def eqimg(slide, tok, key, x, y, mw, mh, align="ctr", valign="ctr", name=None, k=None):
    """One formula or pseudocode block.

    Capped at EQ_K so a short formula is never blown up past the deck's standard
    maths size; shrunk below it only when the box is too small. Pass `k` to match
    the size of a companion formula on the same slide.
    """
    return place(slide, EQ[key], x, y, mw, mh, align, valign, name or key,
                 cap=EQ_K if k is None else k)


ALG_MIN_K = 0.16          # 单张低于这个倍率才考虑切两段（极长清单专用）
# 算法与公式出自同一条渲染管线（同 DPI、同 10pt 基准），所以"一个 PNG 像素 =
# 一个固定物理尺寸"对两者都成立 —— 它们必须用同一个倍率，否则同一份 deck 里
# 算法的字号会和公式不一致。ALG_K 就是给单张路径设的上限。
ALG_K = EQ_K


def alg(slide, tok, key, x, y, w, h, split=True):
    """放一个算法清单（论文三线式，图里已带 "Algorithm N: 名称" 头）。

    近方形的清单塞进 16:9 正文区时只能受高度限制 —— 单张最大也就是 h 那么高，
    字会偏小。若算出来的倍率低于 ALG_MIN_K，且 make_formulas.py 已经切好了左右
    两段，就改成两段并排，把版面横向的余量用上。返回底边 y。
    """
    info = EQ[key]
    single = min(w / info["w"], h / info["h"], ALG_K)
    a, b = key + "_a", key + "_b"
    if split and single < ALG_MIN_K and a in EQ and b in EQ:
        gap = 24
        cw = (w - gap) / 2
        k = min(cw / EQ[a]["w"], h / EQ[a]["h"],
                cw / EQ[b]["w"], h / EQ[b]["h"], ALG_K)
        for col, kk in ((0, a), (1, b)):
            i = EQ[kk]
            _put(slide, i, x + col * (cw + gap) + (cw - i["w"] * k) / 2,
                 y + (h - i["h"] * k) / 2, i["w"] * k, i["h"] * k, kk)
        return y + h
    # 左对齐：算法是"块"，与正文左边界对齐比居中浮动更整齐。cap=ALG_K 防止
    # 短清单被放大到填满栏宽——那会让同一份 deck 里算法的字号互不一致。
    return place(slide, info, x, y, w, h, "l", "ctr", name=key, cap=ALG_K)


def cap(slide, tok, x, y, w, text, labeltxt="图", size=11):
    tb(slide, tok, x, y, w, 18,
       [P([(labeltxt + "　", K.c(tok, "primary"), True),
           (text, K.c(tok, "ink-muted"), False)], size=size, align="ctr", ls=1.2)],
       name="图注")


def eqblock(slide, tok, keys, x, y, w, mh, pad=26, gap=10, name="公式块", k=None):
    """Stack formulas inside one surface-alt block with a primary left bar.

    All of them go in at ONE scale — the point of the block is that the maths
    reads as a single display. If the stack will not fit ``mh`` the whole thing
    is scaled down together, never one formula at a time.
    """
    infos = [EQ[key] for key in keys]
    k = EQ_K if k is None else k
    k = min([k] + [(w - 2 * pad) / i["w"] for i in infos])
    hs = [i["h"] * k for i in infos]
    room = mh - 2 * pad - gap * (len(infos) - 1)
    if sum(hs) > room:
        f = max(0.05, room / max(1e-6, sum(hs)))
        k, hs = k * f, [h * f for h in hs]
    need = 2 * pad + sum(hs) + gap * (len(infos) - 1)
    add(slide, K.rect(K.c(tok, "surface-alt"), K.px(x), K.px(y), K.px(w), K.px(need),
                      name=name + "底"))
    add(slide, K.rect(K.c(tok, "primary"), K.px(x), K.px(y), K.px(3), K.px(need),
                      name=name + "左条"))
    cy = y + pad
    for key, hh in zip(keys, hs):
        iw = EQ[key]["w"] * k
        _put(slide, EQ[key], x + (w - iw) / 2, cy, iw, hh, key)
        cy += hh + gap
    return y + need


def eqgrid(slide, tok, keys, x, y, w, mh, cols=2, gap_x=24, gap_y=10, k=None):
    """Lay N formulas out in `cols` columns at ONE shared scale; returns bottom y.

    Used for the pages whose source is a wall of equations (e.g. 改进后的完整
    模型). A single scale across every column is what keeps the maths consistent.
    """
    per = (len(keys) + cols - 1) // cols
    col_w = (w - gap_x * (cols - 1)) / cols
    s = eqk(keys, w, cols=cols, gap_x=gap_x, gap_y=gap_y, mh=mh) if k is None else k
    for c in range(cols):
        chunk = keys[c * per:(c + 1) * per]
        hsum = sum(EQ[key]["h"] * s for key in chunk) + gap_y * max(0, len(chunk) - 1)
        cx, cy = x + c * (col_w + gap_x), y + (mh - hsum) / 2
        for key in chunk:
            info = EQ[key]
            _put(slide, info, cx, cy, info["w"] * s, info["h"] * s, key)
            cy += info["h"] * s + gap_y
    return y + mh


# ---------------------------------------------------------------------------
# panels
# ---------------------------------------------------------------------------

def callout(slide, tok, x, y, w, h, paragraphs, accent="primary", bg="surface-card"):
    """A tinted panel with a 3px top accent bar — the template's alert/example box."""
    add(slide, K.rect(K.c(tok, accent), K.px(x), K.px(y), K.px(w), K.px(3),
                      name="面板顶条"))
    add(slide, K.rect(K.c(tok, bg), K.px(x), K.px(y + 3), K.px(w), K.px(h - 3),
                      name="面板底"))
    tb(slide, tok, x + 18, y + 14, w - 36, h - 24, paragraphs, name="面板文字")


def card(slide, tok, x, y, w, h, title, paragraphs, accent="accent",
         bg="surface-card", tcol="m"):
    add(slide, K.rect(K.c(tok, accent), K.px(x), K.px(y), K.px(w), K.px(3),
                      name="卡顶"))
    add(slide, K.rect(K.c(tok, bg), K.px(x), K.px(y + 3), K.px(w), K.px(h - 3),
                      name="卡底"))
    if title:
        tb(slide, tok, x + 17, y + 15, w - 34, 26,
           [P(R(tok, [(title, tcol, True)]), size=15)], name="卡标题")
    tb(slide, tok, x + 17, y + (46 if title else 15), w - 34, h - (60 if title else 30),
       paragraphs, name="卡正文")


# ---------------------------------------------------------------------------
# tables (booktabs style: green header, alternating fill, bottom rules only)
# ---------------------------------------------------------------------------

def _no_style(table):
    tbl = table._tbl
    pr = tbl.find(qn("a:tblPr"))
    if pr is None:
        pr = parse_xml('<a:tblPr xmlns:a="http://schemas.openxmlformats.org/'
                       'drawingml/2006/main"/>')
        tbl.insert(0, pr)
    pr.set("firstRow", "0")
    pr.set("bandRow", "0")
    for st in pr.findall(qn("a:tableStyleId")):
        pr.remove(st)
    pr.append(parse_xml('<a:tableStyleId xmlns:a="http://schemas.openxmlformats.org/'
                        'drawingml/2006/main">%s</a:tableStyleId>' % NO_STYLE))


def _cell(tok, cell, text, size, color, bold, fill, align="l", bottom_rule=None):
    cell.fill.solid()
    cell.fill.fore_color.rgb = _rgb(fill)
    cell.margin_left = Emu(9525 * 8)
    cell.margin_right = Emu(9525 * 8)
    cell.margin_top = Emu(9525 * 3)
    cell.margin_bottom = Emu(9525 * 3)
    cell.vertical_anchor = 3  # middle
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = {"l": 1, "c": 2, "r": 3}[align]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    run.font.name = LATIN
    # CJK needs an explicit east-asian typeface, which python-pptx does not set
    rPr = run._r.get_or_add_rPr()
    rPr.append(parse_xml('<a:ea xmlns:a="http://schemas.openxmlformats.org/'
                         'drawingml/2006/main" typeface="%s"/>' % EA))
    pr = cell._tc.find(qn("a:tcPr"))
    if pr is not None:
        pr.set("marL", "76200")
        pr.set("marR", "76200")
    if bottom_rule:
        _bottom_border(cell, bottom_rule)


def _bottom_border(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("a:lnB")):
        tcPr.remove(old)
    ln = parse_xml(
        '<a:lnB xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'w="9525" cap="flat" cmpd="sng" algn="ctr">'
        '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
        '<a:prstDash val="solid"/></a:lnB>' % hex_color)
    # lnB must precede fill properties in a:tcPr
    tcPr.insert(0, ln)


def _rgb(hexstr):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hexstr)


def tbl(slide, tok, heads, rows, widths, x=M, y=CONTENT_Y, w=BODY_W, size=11.5,
        head_size=11, row_h=32, title=None, note=None, emph=(), left=(), ctr=()):
    """A real PowerPoint table. Returns the y just below the table + note.

    ``left``/``ctr`` name the column indices that are NOT right-aligned. Rows in
    ``emph`` are drawn in primary-dark bold (the template's "emphasis row"); a
    single cell can be emphasised the same way by prefixing its text with "*".
    """
    if title:
        _note_fn(slide, tok, x, y, w, title, size=11)
        y += 22
    nrows, ncols = len(rows) + 1, len(heads)
    gf = slide.shapes.add_table(nrows, ncols, K.px(x), K.px(y), K.px(w),
                                K.px(row_h * nrows))
    table = gf.table
    _no_style(table)
    for i, f in enumerate(widths):
        table.columns[i].width = Emu(int(w * f * K.PXE))
    for j, htxt in enumerate(heads):
        al = "l" if j in left else ("c" if j in ctr else "r")
        _cell(tok, table.cell(0, j), htxt, size=head_size, color="FFFFFF",
              bold=True, fill=K.c(tok, "primary"), align=al)
    for i, row in enumerate(rows, start=1):
        hot = (i - 1) in emph
        fill = "FFFFFF" if i % 2 else K.c(tok, "surface-alt")
        for j, val in enumerate(row):
            al = "l" if j in left else ("c" if j in ctr else "r")
            mark = hot or val.startswith("*")
            _cell(tok, table.cell(i, j), val.lstrip("*"), size=size,
                  color=K.c(tok, "primary-dark") if mark else K.c(tok, "ink"),
                  bold=mark, fill=fill, align=al,
                  bottom_rule=K.c(tok, "rule"))
    bottom = y + row_h * nrows
    if note:
        _note_fn(slide, tok, x, bottom + 8, w, note, size=10.5, h=34)
        bottom += 34
    return bottom


# ---------------------------------------------------------------------------
# chrome — lives on the SLIDE, not the layout, so it can be clicked and edited
# ---------------------------------------------------------------------------

EYEBROW_Y = 144


def c_cover(slide, tok):
    # Title and English line are sized to fit ONE line each inside 880px: at the
    # template's 45pt the 16-character title needs 960px, which is both too wide
    # and would run under the cover diamonds (they start at x=972).
    tb(slide, tok, 64, 199, 880, 24,
       [{"runs": [(DECK["cover_kicker"], K.c(tok, "primary"), True)],
         "size": 16, "latin": LATIN, "ea": EA, "spc": 180}], name="封面眉标")
    tb(slide, tok, 64, 236, 880, 72,
       [{"runs": [(DECK["cover_title"], K.c(tok, "ink"), True)],
         "size": 40, "latin": LATIN, "ea": EA, "line_spacing": 1.2}],
       name="封面主标题")
    tb(slide, tok, 64, 322, 880, 32,
       [{"runs": [(DECK["cover_en"], K.c(tok, "ink-muted"), False)],
         "size": 18, "latin": LATIN, "ea": EA}], name="封面副标题")
    ps = [{"runs": [(k + "　　", K.c(tok, "ink-muted"), False),
                    (v, K.c(tok, "ink"), False)],
           "size": 16.5, "latin": LATIN, "ea": EA, "line_spacing": 1.5}
          for k, v in DECK["cover_meta"]]
    tb(slide, tok, 64, 404, 880, 180, ps, name="封面信息")


def c_divider(slide, tok, no):
    tb(slide, tok, 105, 255, 212, 212,
       [{"runs": [("%02d" % no, "FFFFFF", True)], "size": 72,
         "latin": LATIN, "ea": EA, "align": "ctr"}], name="章节数字", anchor="ctr")
    tb(slide, tok, 378, 268, 800, 58,
       [{"runs": [(DECK["sections"][no - 1], K.c(tok, "ink"), True)],
         "size": 36, "latin": LATIN, "ea": EA}], name="章节标题")
    tb(slide, tok, 378, 330, 840, 26,
       [{"runs": [(DECK["sections_en"][no - 1], K.c(tok, "ink-muted"), False)],
         "size": 15, "latin": LATIN, "ea": EA}], name="章节英文")


def c_ending(slide, tok):
    # Text comes from DECK like every other string the author may rewrite.
    tb(slide, tok, 140, 258, 1000, 76,
       [{"runs": [(DECK["end_title"], K.c(tok, "ink"), True)], "size": 52,
         "latin": LATIN, "ea": EA, "align": "ctr"}], name="致谢主标题")
    tb(slide, tok, 140, 344, 1000, 28,
       [{"runs": [(DECK["end_sub"], K.c(tok, "ink-muted"), False)], "size": 22,
         "latin": LATIN, "ea": EA, "align": "ctr", "spc": 300}], name="致谢副标题")
    tb(slide, tok, 140, 404, 1000, 56,
       [{"runs": [(DECK["author_inst"], K.c(tok, "ink-muted"), False)],
         "size": 13, "latin": LATIN, "ea": EA, "align": "ctr",
         "line_spacing": 1.5},
        {"runs": [(DECK["end_date"], K.c(tok, "ink-muted"), False)],
         "size": 13, "latin": LATIN, "ea": EA, "align": "ctr",
         "line_spacing": 1.5}], name="致谢信息")


def eyebrow(slide, tok, no):
    """The '01 · 研究背景' label above the title. On the slide so the author can
    rename the section without opening Slide Master."""
    tb(slide, tok, M, EYEBROW_Y, 900, 22,
       [{"runs": [("%02d · %s" % (no, DECK["sections"][no - 1]),
                   K.c(tok, "primary"), True)],
         "size": 13, "latin": LATIN, "ea": EA}], name="章节眉标")


def c_toc(slide, tok):
    """TOC entries, on the slide so they can be renamed in place."""
    col_w = (BODY_W - 44) / 2
    row_h = 94
    y0 = CONTENT_Y + 14
    for i, (zh, en) in enumerate(zip(DECK["sections"], DECK["sections_en"])):
        col, row = i % 2, i // 2
        x = M + col * (col_w + 44)
        y = y0 + row * row_h
        tb(slide, tok, x, y + 12, 56, 42,
           [{"runs": [("%02d" % (i + 1), K.c(tok, "primary"), True)], "size": 26,
             "latin": LATIN, "ea": EA}], name="目录序号%d" % (i + 1))
        tb(slide, tok, x + 58, y + 12, col_w - 60, 28,
           [{"runs": [(zh, K.c(tok, "ink"), True)], "size": 18,
             "latin": LATIN, "ea": EA}], name="目录标题%d" % (i + 1))
        tb(slide, tok, x + 58, y + 44, col_w - 60, 20,
           [{"runs": [(en, K.c(tok, "ink-muted"), False)], "size": 10,
             "latin": LATIN, "ea": EA, "spc": 40}], name="目录英文%d" % (i + 1))
        add(slide, K.rect(K.c(tok, "rule"), K.px(x), K.px(y + 77),
                          K.px(col_w), K.px(1), name="目录分隔线%d" % (i + 1)))


def gold_sub_bar(slide, tok, section_no, k, n):
    """Gold bar under the current section's segment, length = k/n of the segment."""
    x = M + (section_no - 1) * BAR_PITCH
    w = max(1.0, BAR_W * k / float(n))
    add(slide, K.rect(K.c(tok, "accent"), K.px(x),
                      K.px(BAR_TOP + BAR_H + GOLD_GAP), K.px(w), K.px(GOLD_SUB_H),
                      name="章节内进度%d-%d" % (k, n)))


def _divider(no):
    """Bind the section number so every body fn keeps the simple (slide, tok)."""
    return lambda slide, tok: c_divider(slide, tok, no)


def symlist(slide, tok, x, y, w, h, groups, size=10.5, ls=1.3, name="符号表"):
    """Grouped symbol definitions: a heading, then 'sym  description' lines."""
    ps = []
    for title, items in groups:
        ps.append(P(MT(tok, [(title, "k", True)]), size=size + 0.5, ls=1.25,
                    spc_before=7, spc_after=3))
        for sym, desc in items:
            ps.append(P(mr(sym, _c(tok, "k"), True) + mr("　" + desc, _c(tok, "m"), False),
                        size=size, ls=ls, spc_after=1))
    tb(slide, tok, x, y, w, h, ps, name=name)


# ===========================================================================
# 内容页示例 —— 一个函数演示一种能力，照抄即可
# ===========================================================================

def e_bullets_two_col(slide, tok):
    """双栏项目符号。左栏主色强调、右栏辅色，末尾可加注释。"""
    bullets(slide, tok, M, CONTENT_Y, COL_W, 300, [
        MT(tok, [("数据来源：", "k", True),
                 ("某系统 6 年逐 15 分钟记录，覆盖 237 个观测点。", "m", False)]),
        MT(tok, [("数据质量：", "k", True),
                 ("缺失率 2.3%，样条插值补齐；异常点用 3σ 准则识别。", "m", False)]),
        MT(tok, [("外部变量：", "k", True),
                 ("温度、湿度、辐照度按站点做 IDW 空间插值对齐。", "m", False)]),
    ], name="左栏")
    note(slide, tok, M, CONTENT_Y + 252, COL_W,
         "注：所有序列做 Z-score 标准化，避免量纲影响。")
    bullets(slide, tok, COL2_X, CONTENT_Y, COL_W, 300, [
        MT(tok, [("特征工程：", "n", True),
                 ("滑动窗口长度经网格搜索定为 96 步（24 小时）。", "m", False)]),
        MT(tok, [("日历特征：", "n", True),
                 ("星期、节假日、分时电价区间。", "m", False)]),
        MT(tok, [("数据划分：", "n", True),
                 ("训练 / 验证 / 测试按 7 : 1 : 2 时序切分，不随机打乱。", "m", False)]),
    ], name="右栏")


def e_text_figure(slide, tok):
    """图文：左侧说明文字，右侧等比插图 + 图注。

    img() 自己读 manifest 里的像素宽高来等比缩放，所以不会拉伸变形；
    它返回图片底边 y，交给 cap() 放图注。
    """
    tw = 420
    tb(slide, tok, M, CONTENT_Y, tw, 320, [
        P(MT(tok, [("整体框架分三层：", "m", False), ("感知层", "k", True),
                   ("汇聚量测与外部变量，", "m", False), ("特征层", "k", True),
                   ("提取多尺度模式，", "m", False), ("输出层", "k", True),
                   ("给出带区间的预测结果。", "m", False)]), size=16, ls=1.55),
        P(MT(tok, [("相比点预测，区间预测可直接支撑备用容量决策，这是与既有工作的主要差别。",
                    "m", False)]), size=16, ls=1.55, spc_before=10),
    ], name="正文")
    fx = M + tw + 30
    fw = BODY_W - tw - 30
    # 用横构图的图：竖图在这个宽扁的栏里会被压到只有 1/3 栏宽，看不清（§7.2）。
    b = img(slide, tok, "nc_network", fx, CONTENT_Y - 6, fw, 372)
    cap(slide, tok, fx, b + 8, fw, "系统结构示意（示例插图）", labeltxt="图 1")


def e_cards(slide, tok):
    """三卡并列，顶部 3px 色条。主色/辅色交替靠 accent= 与 bg= 指定。"""
    data = [
        ("数据构建", "整合 6 年逐 15 分钟记录，完成清洗与对齐，形成可复现管线。",
         "accent", "surface-card", "m"),
        ("基线复现", "复现三类基线模型，统一评测协议，便于公平比较。",
         "secondary", "secondary-tint", "m"),
        ("初步结果", "测试集上 MAE 较最优基线下降 8.4%，区间覆盖率接近名义值。",
         "accent", "surface-card", "m"),
    ]
    for i, (title, body, acc, bg, tcol) in enumerate(data):
        card(slide, tok, M + i * (CARD_W + 20), CONTENT_Y, CARD_W, 150, title,
             [P(MT(tok, [(body, "m", False)]), size=13, ls=1.5)],
             accent=acc, bg=bg, tcol=tcol)


def e_table(slide, tok):
    """表格：绿头白字、隔行浅底、只画底边、数字右对齐。

    left= / ctr= 指定**不**右对齐的列（其余列一律右对齐）；emph= 整行强调，
    想让单个单元格强调就在文字前加 "*"（见 tbl 的 docstring）。
    """
    tbl(slide, tok,
        heads=["方法", "MAE", "RMSE", "MAPE / %", "覆盖率 / %"],
        rows=[["基线 A", "0.412", "0.587", "4.71", "—"],
              ["基线 B", "0.386", "0.551", "4.42", "—"],
              ["基线 C", "0.371", "0.534", "4.28", "—"],
              ["本文方法", "0.340", "0.497", "3.95", "94.2"]],
        widths=[0.30, 0.16, 0.16, 0.18, 0.20],
        size=12, head_size=11.5, row_h=30,
        title="表 1　不同方法的误差对比（示例数据，替换为实测结果）",
        note="注：区间预测仅本文方法提供，故基线对应列以「—」表示。",
        left=(0,), emph=(3,))


def e_formula_block(slide, tok):
    """公式块：浅绿底 + 左侧绿条，块内所有公式**共用一个倍率**。

    编号留在页上做文字（可编辑、可搜索），只有公式本体是图片。
    """
    tb(slide, tok, M, CONTENT_Y, BODY_W, 26,
       [P(MT(tok, [("模型对每个分位同时优化，从而直接输出预测区间：", "m", False)]),
          size=16)], name="引语")
    keys = ("demo_objective", "demo_constraint")
    y0 = CONTENT_Y + 40
    b = eqblock(slide, tok, keys, M, y0, BODY_W, 170)
    for key, dy in ((keys[0], 30), (keys[1], 96)):
        if EQ[key].get("tag"):
            tb(slide, tok, M + BODY_W - 78, y0 + dy, 62, 22,
               [P([(EQ[key]["tag"], _c(tok, "u"), False)], size=13, align="r")],
               name="公式编号")
    note(slide, tok, M, b + 10, BODY_W,
         "注：块内两条公式由 eqblock 用同一个倍率放置 —— 见 SKILL §5.3.1。")


def e_formula_grid(slide, tok):
    """公式墙：多列排布，所有列共用一个倍率，翻页时字号不跳。"""
    eqgrid(slide, tok, ("demo_objective", "demo_update",
                        "demo_gap", "demo_subproblem"),
           M, CONTENT_Y, BODY_W, 340, cols=2)
    note(slide, tok, M, CONTENT_Y + 352, BODY_W,
         "注：eqgrid 先算出一个统一倍率再排所有列。某页放不下时应整页统一缩小或拆页，"
         "绝不能只缩其中一张。")


def e_pseudocode(slide, tok):
    """算法清单：论文三线式，图里自带 "Algorithm N: 名称" 头与行号。

    alg() 会在单张放不下（倍率低于 ALG_MIN_K）时自动改用两段并排 —— 见 §5.3.2。
    """
    alg(slide, tok, "demo_pseudo", M, CONTENT_Y, BODY_W, 362)
    note(slide, tok, M, CONTENT_Y + 372, BODY_W,
         "注：算法由 make_formulas.py 用 algorithm2e（ruled 三线式）渲染，"
         "标题与编号自动生成；过长时切两段并排。")


def e_stats(slide, tok):
    """关键数字 + 一张宽图。数字用大号作视觉锚点。

    大号数字的段间距要用 ls≈1.05：段落默认 ls=1.45 会让 50pt 的行高（≈97px）
    超出给它留的框，压到下面的说明文字上。
    """
    stats = [("−8.4%", "MAE 相对最优基线下降", "k"),
             ("94.2%", "95% 名义区间覆盖率", "n"),
             ("6 yr", "训练数据时间跨度", "k")]
    each = BODY_W / 3
    for i, (val, cap_txt, key) in enumerate(stats):
        x = M + i * each
        tb(slide, tok, x, CONTENT_Y - 16, each - 20, 70,
           [P([(val, K.c(tok, "primary" if key == "k" else "secondary"), True)],
              size=48, ls=1.05)], name="数字")
        note(slide, tok, x, CONTENT_Y + 58, each - 20, cap_txt, size=13, h=34)
    # 宽构图的图铺满整宽最清楚；小图并排反而两边都看不清（见 SKILL §7.2）。
    fy = CONTENT_Y + 100
    b = img(slide, tok, "resilience", M, fy, BODY_W, 272)
    cap(slide, tok, M, b + 6, BODY_W, "指标随参数变化（宽构图示例插图）", labeltxt="图 2")


def e_summary(slide, tok):
    """小结：两卡对照（已完成 / 待解决）+ 一条 callout 面板。"""
    h = 146
    card(slide, tok, M, CONTENT_Y, COL_W, h, "已完成", [
        P(MT(tok, [("① 数据管线与评测协议搭建完毕；", "m", False)]), size=13, ls=1.5),
        P(MT(tok, [("② 三类基线复现并统一评测；", "m", False)]), size=13, ls=1.5),
        P(MT(tok, [("③ 损失方案验证有效，覆盖率达标。", "m", False)]), size=13, ls=1.5),
    ], accent="primary", bg="surface-card", tcol="m")
    card(slide, tok, COL2_X, CONTENT_Y, COL_W, h, "待解决", [
        P(MT(tok, [("① 极端样本稀疏，尾部区间偏窄；", "m", False)]), size=13, ls=1.5),
        P(MT(tok, [("② 跨区域迁移时精度衰减明显；", "m", False)]), size=13, ls=1.5),
        P(MT(tok, [("③ 推理时延未满足在线要求。", "m", False)]), size=13, ls=1.5),
    ], accent="alert", bg="surface-alt", tcol="a")
    callout(slide, tok, M, CONTENT_Y + h + 26, BODY_W, 92, [
        P(MT(tok, [("核心问题　", "k", True),
                   ("现有方法在分布偏移下的可靠性如何界定？失效边界在哪里？", "m", False)]),
          size=15, ls=1.5),
    ], accent="primary", bg="surface-card")


def e_references(slide, tok):
    """参考文献列表。

    条目可以来自 prep_refs.py 从源 PDF 抽取的结果（REFS），也可以像这里一样
    直接内联。悬挂缩进用 marL / indent 做。
    """
    items = REFS[:6] if REFS else [
        {"n": 1, "text": "GB/T 7714-2015. 信息与文献 参考文献著录规则[S]. "
                         "北京: 中国标准出版社, 2015."},
        {"n": 2, "text": "Hosseini S, Barker K, Ramirez-Marquez J E. A review of definitions "
                         "and measures of system resilience[J]. Reliability Engineering & "
                         "System Safety, 2016, 145: 47-61."},
        {"n": 3, "text": "Wang Y, Chen Q, Hong T, et al. Review of smart meter data analytics[J]. "
                         "IEEE Transactions on Smart Grid, 2019, 10(3): 3125-3148."},
        {"n": 4, "text": "示例条目：换成你自己的文献，或用 prep_refs.py 从源 PDF 自动抽取。"},
    ]
    ps = [P([("[%d] " % it["n"], K.c(tok, "primary"), True),
             (it["text"], K.c(tok, "ink"), False)],
            size=12, ls=1.5, spc_after=12,
            marL=285750, indent=-285750) for it in items]
    tb(slide, tok, M, CONTENT_Y, BODY_W, 340, ps, name="参考文献")


SLIDE_PLAN = [
    # (版式名, 标题占位符文字, 眉标章节号 或 None, 正文函数 或 None)
    ("封面", None, None, c_cover),
    ("目录", "目录", None, c_toc),

    ("章节过渡 01", None, None, _divider(1)),
    ("内容 01", "数据来源与预处理", 1, e_bullets_two_col),       # 1/2
    ("内容 01", "模型总体框架", 1, e_text_figure),               # 2/2

    ("章节过渡 02", None, None, _divider(2)),
    ("内容 02", "三项已完成工作", 2, e_cards),                   # 1/2
    ("内容 02", "不同方法的精度对比", 2, e_table),                # 2/2

    ("章节过渡 03", None, None, _divider(3)),
    ("内容 03", "目标函数与约束", 3, e_formula_block),            # 1/3
    ("内容 03", "完整模型（公式墙）", 3, e_formula_grid),          # 2/3
    ("内容 03", "求解流程", 3, e_pseudocode),                     # 3/3

    ("章节过渡 04", None, None, _divider(4)),
    ("内容 04", "精度提升与消融", 4, e_stats),                    # 1/2
    ("内容 04", "阶段性总结", 4, e_summary),                      # 2/2

    ("章节过渡 05", None, None, _divider(5)),
    ("内容 05", "参考文献", 5, e_references),                     # 1/1

    ("致谢", None, None, c_ending),
]


def build_slides(prs, tok):
    by_name = {l.name: l for l in prs.slide_masters[0].slide_layouts}
    # Pages per section, so the gold sub-bar can be k/n of the segment.
    from collections import Counter
    per_section = Counter(no for _, _, no, _ in SLIDE_PLAN if no is not None)
    seen = Counter()

    made = []
    for layout_name, title, section_no, body_fn in SLIDE_PLAN:
        slide = prs.slides.add_slide(by_name[layout_name])
        if title:
            # Placeholders clone empty from the layout, which is correct for a
            # template but looks unfinished in the worked examples.
            for ph in slide.placeholders:
                if ph.placeholder_format.type == PP_PLACEHOLDER.TITLE:
                    ph.text = title
                    break
        if section_no is not None:
            eyebrow(slide, tok, section_no)
            seen[section_no] += 1
            gold_sub_bar(slide, tok, section_no, seen[section_no],
                         per_section[section_no])
        if body_fn is not None:
            body_fn(slide, tok)
        made.append(slide)
    return made

